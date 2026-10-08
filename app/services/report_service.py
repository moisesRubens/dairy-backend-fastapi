from datetime import date, datetime, time, timedelta
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy.orm import selectinload
from database.model import Order, OrderSalePoint, SalePoints
from database.money import money, subtotal
from decimal import Decimal


def report_period(day: date, period: str):
    start = day if period == 'day' else day - timedelta(days=day.weekday())
    end = start + timedelta(days=1 if period == 'day' else 7)
    return start, end


def chart_report(session, day: date, period: str):
    start, end = report_period(day, period)
    count = 24 if period == 'day' else 7
    points = {p.id: {'id': str(p.id), 'name': p.name, 'values': [money(0)] * count}
              for p in session.query(SalePoints).order_by(SalePoints.name).all()}
    products = {}
    orders = session.query(OrderSalePoint).join(Order).filter(
        Order.order_date >= datetime.combine(start, time.min),
        Order.order_date < datetime.combine(end, time.min),
    ).options(selectinload(OrderSalePoint.order).selectinload(Order.item_order)).all()
    seen = set()
    for link in orders:
        order = link.order
        bucket = order.order_date.hour if period == 'day' else (order.order_date.date() - start).days
        points[link.sale_point_id]['values'][bucket] += order.total_value
        if order.id in seen:
            continue
        seen.add(order.id)
        for item in order.item_order:
            key = str(item.product_id) if item.product_id is not None else 'deleted:' + item.product_name
            series = products.setdefault(key, {'id': key, 'name': item.product_name, 'values': [money(0)] * count})
            quantity = next((getattr(item, unit) for unit in ('amount', 'kg', 'liters')
                             if getattr(item, unit) is not None and getattr(item, unit) >= 0), 0)
            series['values'][bucket] += subtotal(item.item_price, quantity)
    for series in [*points.values(), *products.values()]:
        series['values'] = [float(money(value)) for value in series['values']]
    labels = [f'{hour:02d}h' for hour in range(24)] if period == 'day' else [
        (start + timedelta(days=index)).strftime('%d/%m') for index in range(7)]
    return {'labels': labels, 'sale_points': list(points.values()),
            'products': sorted(products.values(), key=lambda series: series['name'])}


def export_report(session, day: date, period: str):
    start, end = report_period(day, period)
    lower, upper = datetime.combine(start, time.min), datetime.combine(end, time.min)
    points = {p.id: p.name for p in session.query(SalePoints).order_by(SalePoints.id).all()}
    orders = session.query(OrderSalePoint).join(Order).filter(
        Order.order_date >= lower, Order.order_date < upper,
    ).options(selectinload(OrderSalePoint.order).selectinload(Order.item_order)).order_by(
        OrderSalePoint.sale_point_id, Order.order_date, Order.id,
    ).all()
    wb = Workbook()
    sales = wb.active
    sales.title = 'Pedidos'
    sales.append(['Data de in\u00edcio', 'Nome do ponto de venda', 'ID do pedido', 'Valor do pedido', 'Status do pedido', 'Nome do item', 'Valor do item'])
    for link in orders:
        order = link.order
        for item in sorted(order.item_order, key=lambda item: item.product_id):
            quantity = next((getattr(item, unit) for unit in ('amount', 'kg', 'liters')
                             if getattr(item, unit) is not None and getattr(item, unit) >= 0), 0)
            sales.append([start, points.get(link.sale_point_id, ''), order.id,
                          order.total_value, order.order_status, item.product_name,
                          subtotal(item.item_price, quantity)])
    for sheet in wb:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='333333')
            sheet.column_dimensions[cell.column_letter].width = 22
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                # User-provided names/descriptions are text, never Excel formulas.
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                elif isinstance(cell.value, datetime):
                    cell.value = cell.value.replace(tzinfo=None)
                    cell.number_format = 'dd/mm/yyyy hh:mm'
                elif isinstance(cell.value, date):
                    cell.number_format = 'dd/mm/yyyy'
                elif isinstance(cell.value, (float, Decimal)):
                    cell.number_format = '0.00'
    result = BytesIO()
    wb.save(result)
    return result.getvalue(), f'vendas_{start}_{end - timedelta(days=1)}.xlsx'
