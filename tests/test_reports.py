import unittest
from datetime import date, datetime
from io import BytesIO
from openpyxl import load_workbook
import test_sale_points as fixtures
from database.model import SalePoints, Product, Order, ItemsOrder, OrderSalePoint, RetiradaProduto
from services.report_service import report_period


class ReportTests(unittest.TestCase):
    setUp = fixtures.SalePointIntegrationTests.setUp
    tearDown = fixtures.SalePointIntegrationTests.tearDown
    create = fixtures.SalePointIntegrationTests.create

    def seed(self):
        actor = self.create('Admin').json()['id']
        seller = self.create('Vendedor').json()['id']
        with self.sessions() as session:
            session.get(SalePoints, actor).level = 1
            product = Product(name='Queijo atual', price=50, kg=10)
            session.add(product)
            session.flush()
            for point, day, total, paid in [(actor, 1, 20, True), (seller, 4, 30, False), (seller, 5, 40, True), (seller, 28, 50, True)]:
                order = Order(order_date=datetime(2026, 10 if day < 28 else 9, day, 23, 59), total_value=total, status=paid, discount_value=0, description='=SUM(A1)')
                session.add(order)
                session.flush()
                session.add(OrderSalePoint(order_id=order.id, sale_point_id=point))
                session.add(ItemsOrder(order_id=order.id, product_id=product.id, product_name='Queijo histórico', item_price=40, kg=total / 40))
            for point, day, active in [(actor, 1, True), (seller, 4, False), (seller, 5, True)]:
                session.add(RetiradaProduto(sale_point_id=point, product_id=product.id, data=datetime(2026, 10, day), unidade='kg', taken_quantity=2, sold_quantity=1, remaining_quantity=1 if active else 0, total_value=40, status=active))
            session.commit()

    def workbook(self, period, day='2026-10-01'):
        response = self.client.get('/sale-points/reports/export', params={'day': day, 'period': period})
        self.assertEqual(response.status_code, 200, response.text if response.status_code != 200 else '')
        self.assertIn('.xlsx', response.headers['content-disposition'])
        self.assertIn('spreadsheetml', response.headers['content-type'])
        return load_workbook(BytesIO(response.content))

    def test_charts_bucket_sales_and_products_without_repeating_order_total(self):
        self.seed()
        with self.sessions() as session:
            session.add(ItemsOrder(order_id=1, product_name='Outro item', item_price=5, amount=2))
            session.commit()
        response = self.client.get('/sale-points/reports/charts?day=2026-10-01&period=week')
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data['labels'], ['28/09', '29/09', '30/09', '01/10', '02/10', '03/10', '04/10'])
        points = {s['name']: s['values'] for s in data['sale_points']}
        self.assertEqual(points['Admin'], [0, 0, 0, 20, 0, 0, 0])
        self.assertEqual(points['Vendedor'], [50, 0, 0, 0, 0, 0, 30])
        products = {s['name']: s['values'] for s in data['products']}
        self.assertEqual(products['Outro item'][3], 10)
        day = self.client.get('/sale-points/reports/charts?day=2026-10-01').json()
        self.assertEqual(len(day['labels']), 24)
        self.assertEqual(sum(sum(s['values']) for s in day['sale_points']), 20)
        self.assertEqual(sum(s['values'][23] for s in day['sale_points']), 20)

    def test_charts_require_admin(self):
        self.create()
        self.assertEqual(self.client.get('/sale-points/reports/charts?day=2026-10-01').status_code, 403)

    def test_day_has_exact_columns_and_historical_item_subtotal(self):
        self.seed()
        wb = self.workbook('day')
        self.assertEqual(wb.sheetnames, ['Pedidos'])
        sheet = wb['Pedidos']
        self.assertEqual([cell.value for cell in sheet[1]], [
            'Data de in\u00edcio', 'Nome do ponto de venda', 'ID do pedido',
            'Valor do pedido', 'Status do pedido', 'Nome do item', 'Valor do item'])
        self.assertEqual(sheet.max_column, 7)
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet['A2'].value.date(), date(2026, 10, 1))
        self.assertEqual(sheet['B2'].value, 'Admin')
        self.assertEqual(sheet['D2'].value, 20)
        self.assertEqual(sheet['E2'].value, 'pago')
        self.assertEqual(sheet['F2'].value, 'Queijo hist\u00f3rico')
        self.assertEqual(sheet['G2'].value, 20)
        self.assertEqual(sheet.freeze_panes, 'A2')

    def test_week_filters_dates_and_repeats_order_fields_per_item(self):
        self.seed()
        with self.sessions() as session:
            product = Product(name='Outro', price=5, amount=10)
            session.add(product)
            session.flush()
            session.add(ItemsOrder(order_id=1, product_id=product.id,
                                  product_name='=SUM(A1)', item_price=5, amount=2))
            session.get(Order, 1).discount_value = 2
            session.get(Order, 1).total_value = 28
            session.commit()
        sheet = self.workbook('week')['Pedidos']
        self.assertEqual(sheet.max_row, 5)  # Sep 28 through Oct 4, two items in order 1
        self.assertEqual(sheet['A2'].value.date(), date(2026, 9, 28))
        self.assertEqual(sheet['C2'].value, sheet['C3'].value)
        self.assertEqual(sheet['D2'].value, 28)
        self.assertEqual(sheet['D3'].value, 28)
        self.assertEqual(sheet['E2'].value, 'desconto')
        self.assertEqual(sheet['F3'].value, '=SUM(A1)')
        self.assertEqual(sheet['F3'].data_type, 's')
        self.assertEqual(sheet['G3'].value, 10)
        self.assertEqual(sheet['E5'].value, 'pendente')
        self.assertEqual(report_period(date(2026, 10, 4), 'week'), (date(2026, 9, 28), date(2026, 10, 5)))

    def test_empty_export_still_has_headers_and_admin_authorization(self):
        id = self.create().json()['id']
        self.assertEqual(self.client.get('/sale-points/reports/export?day=2026-10-01').status_code, 403)
        with self.sessions() as session:
            session.get(SalePoints, id).level = 1
            session.commit()
        wb = self.workbook('day')
        self.assertEqual(wb['Pedidos'].max_row, 1)
        self.assertEqual(wb['Pedidos'].max_column, 7)
        self.assertEqual(self.client.get('/sale-points/reports/export?day=invalid').status_code, 422)
        self.assertEqual(self.client.get('/sale-points/reports/export?day=2026-10-01&period=month').status_code, 422)
